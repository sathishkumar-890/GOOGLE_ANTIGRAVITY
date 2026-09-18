import openpyxl

wb_fun = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws_fun = wb_fun['Allocation Sheet']

entries = []
for r in range(12, ws_fun.max_row + 1):
    func = ws_fun.cell(r, 9).value
    mach = ws_fun.cell(r, 10).value
    hyper = ws_fun.cell(r, 3).value
    if mach and mach != 'Folder' and func and func != 'Folder':
        entries.append((str(mach).strip(), str(func).strip(), str(hyper).strip() if hyper else ''))

# Search for winch, windlass, iccp, steering, loading, pump, etc.
keywords = ['WINCH', 'WINDLASS', 'ICCP', 'STEERING', 'LAODING', 'LOADING', 'PUMP', 'FAN', 'COMPRESSOR', 'PURIFIER', 'HEATER', 'COOLER']

for kw in keywords:
    matches = [e for e in entries if kw in e[0].upper() or kw in e[1].upper()]
    print(f"\n--- Keyword '{kw}' matches ({len(matches)}): ---")
    seen = set()
    for m, f, h in matches:
        if (m, f) not in seen:
            seen.add((m, f))
            print(f"  Mach: {m:40s} | Func: {f}")
