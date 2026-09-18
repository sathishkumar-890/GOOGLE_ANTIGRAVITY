import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import shutil
import os

EXCEL_PATH = '01_VL_EXTRACTION/VL_IRENES REWARD.xlsx'
ROOT_EXCEL_PATH = 'VL_IRENES REWARD.xlsx'
BACKUP_PATH = '01_VL_EXTRACTION/VL_IRENES REWARD_BACKUP_BEFORE_COLS.xlsx'

shutil.copy2(EXCEL_PATH, BACKUP_PATH)
print(f'Backup saved to {BACKUP_PATH}')

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
ws = wb.active

# 1. Revert the 6 duplicate model rows to 'NA' and remove green fill
duplicate_model_rows = [312, 313, 399, 420, 441, 490]
no_fill = PatternFill(fill_type=None)

for r in duplicate_model_rows:
    cell = ws.cell(row=r, column=8)
    cell.value = 'NA'
    cell.fill = no_fill
    print(f'Reverted Row {r} to NA (was duplicate model)')

# 2. Insert 2 columns at index 9 (after Column 8 S.NO)
ws.insert_cols(9, 2)
print('Inserted 2 columns at index 9')

from copy import copy

# Header formatting
header_fill = copy(ws.cell(row=1, column=8).fill)
header_font = Font(name='Arial', size=10, bold=True, color='000000')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='FF000000'),
    right=Side(style='thin', color='FF000000'),
    top=Side(style='thin', color='FF000000'),
    bottom=Side(style='thin', color='FF000000')
)

cell_h9 = ws.cell(row=1, column=9)
cell_h9.value = 'S.NO PDF NAME'
cell_h9.font = header_font
cell_h9.alignment = header_align
cell_h9.border = thin_border
if header_fill:
    cell_h9.fill = header_fill

cell_h10 = ws.cell(row=1, column=10)
cell_h10.value = 'S.NO PAGE NO'
cell_h10.font = header_font
cell_h10.alignment = header_align
cell_h10.border = thin_border
if header_fill:
    cell_h10.fill = copy(header_fill)

# 3. Update Column E formulas from referencing L to N
for r in range(2, ws.max_row + 1):
    cell_e = ws.cell(row=r, column=5)
    val = str(cell_e.value or '')
    if f'L{r}' in val:
        cell_e.value = val.replace(f'L{r}', f'N{r}')

# 4. Map of row -> (pdf_name, page_no)
pdf_page_map = {
    # Main Engine
    268: ('FINAL DRAWINGS.PDF', '5'),
    # D/G Engines
    92: ('VOL 2.pdf', '643'),
    93: ('VOL 2.pdf', '644'),
    94: ('VOL 2.pdf', '645'),
    # Emergency Generator Set
    124: ('1-1. Specification for H4235.pdf', '1'),
    111: ('1-1. Specification for H4235.pdf', '2'),
    # Kangrim Composite Boiler
    70: ('M-62 COMPOSITE BOILER.pdf', '1'),
    # Alfa Laval Plate Heat Exchangers
    60: ('M-55 PLATE TYPE HEAT EXCHANGER.pdf', '2'),
    61: ('M-55 PLATE TYPE HEAT EXCHANGER.pdf', '2'),
    72: ('M-55 PLATE TYPE HEAT EXCHANGER.pdf', '3'),
    264: ('M-55 PLATE TYPE HEAT EXCHANGER.pdf', '3'),
    426: ('M-55 PLATE TYPE HEAT EXCHANGER.pdf', '2'),
    439: ('M-55 PLATE TYPE HEAT EXCHANGER.pdf', '3'),
    # Kumkang Air Reservoirs
    48: ('M-49 AIR RESERVOIR.pdf', '1'),
    49: ('M-49 AIR RESERVOIR.pdf', '1'),
    332: ('M-49 AIR RESERVOIR.pdf', '1'),
    401: ('M-49 AIR RESERVOIR.pdf', '1'),
    # Mytec S&T Heat Exchangers
    287: ('M-56 S & T TYPE HEAT EXCHANGER.pdf', '1'),
    410: ('M-56 S & T TYPE HEAT EXCHANGER.pdf', '1'),
    # Samkun Century
    12: ('M-50 AIR DRYER.pdf', '1'),
    418: ('M-54 F.W. SUPPLY UNIT.pdf', '1'),
    # Hyundai Elevator
    82: ('A-20 ELEVATOR.pdf', '1'),
    # HGS Consoles & Gyro
    41: ('240604 HMD 4193 BCC F 01 (FINAL).pdf', '1'),
    112: ('20240510_HM-H4193-94_4232-35 ECC WD H03 (FULL).pdf', '1'),
    38: ('240604 HMD 4193 BCC F 01 (FINAL).pdf', '1'),
    23: ('R2D2022_FINAL.pdf', '1'),
    170: ('R2D2022_FINAL.pdf', '1'),
    # Kongsberg AMS
    129: ('E-53,58 ER ALARM AND MONITORING SYSTEM MIMIC.pdf', '1'),
    # Trelleborg Shaft Power Meter
    336: ('E-42 ME HORSE POWER METER.pdf', '1'),
    # Daeyang Anemometer
    18: ('E-37 ANEMOMETER & ANEMOSCOPE.pdf', '1'),
    # Saracom Magnetic Compass
    260: ('MAGNETIC COMPASS.pdf', '1'),
    # HME Electrical & Comm
    42: ('E-68 BATTERY CHARGER & DIST. BOARD.PDF', '1'),
    427: ('E-47 MISCELLANEOUS ALARM SYS.pdf', '1'),
    480: ('E-47 MISCELLANEOUS ALARM SYS.pdf', '1'),
    484: ('E-50 LIGHT SIGNAL COLUMN (INCL. ER DEAD-MAN ALARM).pdf', '1'),
    # Dongyang Panels
    422: ('E-73 WH GROUP CONTROL PANEL.pdf', '1'),
    495: ('E-46 WHEEL HOUSE GAUGE BOARD.pdf', '1'),
    # Daeyang Reefer Sockets & Monitoring
    478: ('CCT-HMD H4235 RD9 RCMS Final-R0.pdf', '1'),
    479: ('E-76 RECEPTACLE FOR REEFER CONTAINER.pdf', '1'),
    # Sehong Systems
    22: ('M-53 M.G.P.S..PDF', '1'),
    334: ('M-40 SHAFT EARTHING DEVICE.PDF', '7'),
    428: ('H-25_I.C.C.P SYSTEM.pdf', '1'),
    # Daelim Loose FFA
    254: ('P-20 LOOSE FIRE FIGHTING APPLIANCE.pdf', '1'),
    # Jung-A Marine Ladders
    16: ('F-34 ACCOMMODATION LADDER.pdf', '1'),
    17: ('F-34 ACCOMMODATION LADDER.pdf', '1'),
    472: ('F-35 PILOT SLANT LADDER.pdf', '1'),
    473: ('F-35 PILOT SLANT LADDER.pdf', '1'),
    # Sangsangin Crane
    83: ('M-60 ER OVERHEAD CRANE.pdf', '1'),
    # HMMCO Incinerator
    429: ('M-61 INCINERATOR.pdf', '1'),
    # MNSI FO Supply Unit
    417: ('Final Drawing for 4235 F.O Supply Unit for ME & GE_240701.pdf', '1'),
    # Anchors & Chain
    19: ('F-29 ANCHOR CHAIN.pdf', '2'),
    20: ('F-28 ANCHOR.pdf', '2'),
    21: ('F-28 ANCHOR.pdf', '2'),
    # Wonkwang Emergency Shut-off Valves
    125: ('M-31 EMERGENCY SHUT-OFF VALVE.pdf', '4'),
    # ABB Energy Saving System
    126: ('3BKA3266897 HMD H4235 ABB EES Final Document RF.pdf', '1'),
    # Fain CO2 System
    145: ('P-19 CO2 EXTINGUISHING SYSTEM.pdf', '3'),
    # SKF Coupling Bolts
    219: ('14714.pdf', '1'),
    220: ('14715.pdf', '1'),
    # Safety Appliances
    242: ('F-37 LIFE BOAT DAVIT.pdf', '7'),
    245: ('P-22 ELEC. PNEUMATIC TYPE TANK LEVEL GAUGE.pdf', '3'),
    248: ('F-36 LIFE BOAT.pdf', '7'),
    256: ('F-38 RESCUE BOAT.pdf', '7'),
    436: ('F-40 LIFE RAFT.pdf', '8'),
    437: ('F-40 LIFE RAFT.pdf', '10'),
    438: ('F-40 LIFE RAFT.pdf', '12'),
    # MRC Comm Systems
    69: ('E-48,49,57,64 MRC.pdf', '1'),
    266: ('E-48,49,57,64 MRC.pdf', '1'),
    301: ('E-48,49,57,64 MRC.pdf', '1'),
    342: ('E-48,49,57,64 MRC.pdf', '1'),
    378: ('E-48,49,57,64 MRC.pdf', '1'),
    483: ('E-48,49,57,64 MRC.pdf', '1'),
    491: ('E-48,49,57,64 MRC.pdf', '1'),
    # JRC Navigation Suite (Order 2DJ-3665)
    14: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    39: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    73: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    113: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    114: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    135: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    235: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    286: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    290: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    291: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    293: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    294: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    318: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    319: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    340: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    361: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    362: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    363: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    364: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    372: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    373: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    482: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    494: ('E-33,34~36,39,40,44,45,62,63,84 JRC.pdf', '1'),
    # KEMEL Bearings & Seals
    237: ('M-39 STERN TUBE SEAL & BEARINGS.pdf', '1'),
    348: ('M-39 STERN TUBE SEAL & BEARINGS.pdf', '1'),
    349: ('M-39 STERN TUBE SEAL & BEARINGS.pdf', '1'),
    486: ('M-39 STERN TUBE SEAL & BEARINGS.pdf', '1'),
    487: ('M-39 STERN TUBE SEAL & BEARINGS.pdf', '1'),
    # Jung Gong CVS & Wipers
    408: ('A-23 WINDOW WIPER & CVS.pdf', '1'),
    409: ('A-23 WINDOW WIPER & CVS.pdf', '1'),
    496: ('A-23 WINDOW WIPER & CVS.pdf', '1'),
    497: ('A-23 WINDOW WIPER & CVS.pdf', '1'),
    498: ('A-23 WINDOW WIPER & CVS.pdf', '1'),
    # Shin Myung Tech Davits
    415: ('F-31 F.O HOSE DAVIT.pdf', '1'),
    416: ('F-31 F.O HOSE DAVIT.pdf', '1'),
    # Flutek Winches & Windlass
    391: ('F-30 DECK MACHINERY.pdf', '1'),
    392: ('F-30 DECK MACHINERY.pdf', '1'),
    393: ('F-30 DECK MACHINERY.pdf', '1'),
    394: ('F-30 DECK MACHINERY.pdf', '1'),
    397: ('F-30 DECK MACHINERY.pdf', '1'),
    398: ('F-30 DECK MACHINERY.pdf', '1'),
    # MacGregor Hatch Cover
    407: ('HMD 4235_Hatch cover_Final.pdf', '1'),
    # Tanktech Anti-Heeling Pump
    400: ('P-21 ANTI-HEELING PUMP.pdf', '1'),
    # Hoppe Korea VRCS
    380: ('P-17 VALVE REMOTE CONTROL SYSTEM.pdf', '1'),
    # Hi-Air ECR AC
    115: ('M-28 PACKAGED AIR CON. FOR E.C.R.pdf', '1'),
    # Hi-Air Ventilation Fans
    84: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    171: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    296: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    371: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    381: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    412: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    492: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    493: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    4: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    5: ('A-21 ACCOMMODATION VENTILATION.pdf', '1'),
    # Taiko Pumps
    32: ('H4235 INSPECTION REPORT.PDF', '1'),
    34: ('H4235 INSPECTION REPORT.PDF', '1'),
    35: ('H4235 INSPECTION REPORT.PDF', '1'),
    36: ('H4235 INSPECTION REPORT.PDF', '1'),
    37: ('H4235 INSPECTION REPORT.PDF', '1'),
    86: ('H4235 INSPECTION REPORT.PDF', '1'),
    146: ('H4235 INSPECTION REPORT_EMERGENCY FIRE PUMP.PDF', '1'),
    210: ('H4235 INSPECTION REPORT.PDF', '1'),
    212: ('H4235 INSPECTION REPORT.PDF', '1'),
    252: ('H4235 INSPECTION REPORT.PDF', '1'),
    253: ('H4235 INSPECTION REPORT.PDF', '1'),
    255: ('H4235 INSPECTION REPORT.PDF', '1'),
    257: ('H4235 INSPECTION REPORT.PDF', '1'),
    258: ('H4235 INSPECTION REPORT.PDF', '1'),
    259: ('H4235 INSPECTION REPORT.PDF', '1'),
    267: ('H4235 INSPECTION REPORT.PDF', '1'),
    276: ('H4235 INSPECTION REPORT.PDF', '1'),
    282: ('H4235 INSPECTION REPORT.PDF', '1'),
    283: ('H4235 INSPECTION REPORT.PDF', '1'),
    305: ('H4235 INSPECTION REPORT.PDF', '1'),
    306: ('H4235 INSPECTION REPORT.PDF', '1'),
    341: ('H4235 INSPECTION REPORT.PDF', '1'),
    343: ('H4235 INSPECTION REPORT.PDF', '1'),
    419: ('H4235 INSPECTION REPORT.PDF', '1'),
    423: ('H4235 INSPECTION REPORT.PDF', '1'),
    424: ('H4235 INSPECTION REPORT.PDF', '1'),
    425: ('H4235 INSPECTION REPORT.PDF', '1'),
    440: ('H4235 INSPECTION REPORT.PDF', '1'),
    488: ('H4235 INSPECTION REPORT.PDF', '1'),
    489: ('H4235 INSPECTION REPORT.PDF', '1'),
    # Phase 1 machinery
    2: ('1. HMD 4235 FINAL DWG (OBS).pdf', '1'),
    44: ('(FINAL DWG)HMD_4193_94_4232_33_34_35 ECS-600Bx1_REV.2_240528.pdf', '1'),
    46: ('H4235 FINAL.pdf', '1'),
    47: ('H4235 FINAL.pdf', '1'),
    104: ('H4235 DRILLING MACHINE FINAL DRAWING & INSTRUCTION MANUAL.pdf', '1'),
    122: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    123: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    151: ('1. Final_Drawings.pdf', '1'),
    209: ('1. Final_Drawings.pdf', '1'),
    244: ('M-58 DRILLING & GRINDING MACHINE.pdf', '1'),
    262: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    263: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    279: ('1. Final_Drawings.pdf', '1'),
    280: ('1. Final_Drawings.pdf', '1'),
    335: ('F-26 STEERING GEAR.pdf', '1'),
    355: ('H4235 GRINDING MACHINE FINAL  DRAWING & INSTRUCTION MANUAL.pdf', '1'),
    387: ('E-74 WHISTLE.pdf', '1'),
    390: ('M-52 SEWAGE TREATMENT PLANT.pdf', '1'),
    402: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    403: ('F-27 BOW(TUNNEL) THRUSTER.pdf', '1'),
    405: ('F-27 BOW(TUNNEL) THRUSTER.pdf', '1'),
    414: ('2.Final Drawing.pdf', '1'),
    431: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    432: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    433: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    434: ('230619 KD222001 For Final DWG (For HMD 4193s TR.).pdf', '1'),
    442: ('F-32 MONORAIL HOIST CRANE.pdf', '1'),
    481: ('F-39 RESCUE BOAT DAVIT.pdf', '1'),
    499: ('H4235 FINAL R2.pdf', '1'),
    500: ('AMP SYSTEM.pdf', '1'),
    95: ('E-77 DIESEL GENERATOR.pdf', '1'),
    96: ('E-77 DIESEL GENERATOR.pdf', '1'),
    97: ('E-77 DIESEL GENERATOR.pdf', '1'),
}

for r in [130, 131, 132, 133, 307, 384]:
    pdf_page_map[r] = ('240708_hmd04235fm_final dwg.pdf', '1')

for r in [370, 406, 471, 485] + list(range(443, 471)):
    pdf_page_map[r] = ('240708_hmd04235fh_final dwg.pdf', '1')

data_font = Font(name='Calibri', size=11)
left_align = Alignment(horizontal='left', vertical='center')
center_align = Alignment(horizontal='center', vertical='center')

populated = 0
for r in range(2, ws.max_row + 1):
    c9 = ws.cell(row=r, column=9)
    c10 = ws.cell(row=r, column=10)
    c9.border = thin_border
    c10.border = thin_border
    c9.font = data_font
    c10.font = data_font
    c9.alignment = left_align
    c10.alignment = center_align

    cell_h = ws.cell(row=r, column=8)
    color_h = getattr(cell_h.fill.start_color, 'rgb', None) if cell_h.fill else None
    if color_h in ['FF92D050', '92D050']:
        mapping = pdf_page_map.get(r)
        if mapping:
            pdf_name, page_no = mapping
            c9.value = pdf_name
            c10.value = page_no
            populated += 1
        else:
            c9.value = 'MANUALS_REWARD'
            c10.value = '1'
            populated += 1

print(f'Populated {populated} rows with PDF Name and Page No')

ws.column_dimensions['I'].width = 38
ws.column_dimensions['J'].width = 15

wb.save(EXCEL_PATH)
print(f'Saved updated workbook to {EXCEL_PATH}')

shutil.copy2(EXCEL_PATH, ROOT_EXCEL_PATH)
print(f'Synchronized root replica to {ROOT_EXCEL_PATH}')
