import openpyxl, re

wb_main = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC.xlsx', data_only=True)
ws_main = wb_main.active

wb_fun = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES WISDOM_FUN.xlsx', data_only=True)
ws_fun = wb_fun.active

# Extract all (MACHINERY NAME, FUNCTION, MANUAL HYPERLINK) from FUN
fun_entries = []
for r in range(7, ws_fun.max_row + 1):
    func = ws_fun.cell(r, 13).value
    mach = ws_fun.cell(r, 14).value
    man = ws_fun.cell(r, 3).value
    if func and func != 'Folder' and mach and mach != 'Folder':
        fun_entries.append((str(mach).strip(), str(func).strip(), str(man).strip() if man else ''))

# Build a lookup function
def get_function_for_row(row_idx, loc_code, loc_desc, sys_desc, maker, model):
    ld = loc_desc.upper().strip()
    sd = sys_desc.upper().strip()
    lc = loc_code.upper().strip()
    
    # Specific rules matching C.V. IRENES WISDOM_FUN.xlsx
    
    # 1. Loading Computer
    if 'LOADING COMPUTER' in ld or 'LCMP' in lc:
        return 'ELECTRONIC EQUIPMENT', 'LAODING COMPUTER'
        
    # 2. ICCP
    if 'I.C.C.P' in ld or 'ICCP' in ld or 'ICCP' in lc:
        return 'ICCP SYSTEM', 'ICCP'
        
    # 3. Steering Gear
    if ('STEERING GEAR' in ld or 'SGR' in lc) and 'FAN' not in ld:
        return 'STEERING SYSTEM', 'STEERING GEAR'
        
    # 4. Anchor & Chain
    if 'ANCHOR CHAIN' in ld or 'ANCCHAIN' in lc:
        return 'ANCHOR CHAIN', 'ANCHOR CHAIN'
    if 'ANCHOR' in ld or 'ANCHOR' in lc:
        return 'ANCHOR&MOORING WINDLASS', 'ANCHOR'
        
    # 5. Mooring Winch & Windlass
    if 'MOORING WINCH' in ld or 'YMRWC' in lc:
        return 'MOORING WINCH', 'MOORING WINCH'
    if 'WINDLASS' in ld or 'YWDLS' in lc:
        return 'WINDLASS', 'WINDLASS'
        
    # 6. Bow Thruster
    if 'BOW THRUSTER' in ld or 'BOWTH' in lc:
        if 'TRANSFORMER' in ld:
            return 'ELECTRICAL SYSTEM', 'POWER & LIGHTING TRANSFORMER'
        if 'PANEL' in ld or 'STARTER' in ld:
            return 'PROPELLER, THRUSTER, STERN TUBE, SHAFTING', 'BOW THRUSTER STARTER'
        if 'FAN' in ld:
            return 'HEATING, VENTILATION & AIR CONDITIONING', 'VENTILATION FAN'
        return 'PROPELLER, THRUSTER, STERN TUBE, SHAFTING', 'BOW THRUSTER'
        
    # 7. Davits & Cranes
    if 'F.O HOSE DAVIT' in ld or 'FOHDAV' in lc:
        return 'DAVIT', 'F.O HOSE DAVIT'
    if 'MONORAIL' in ld or 'MONHCR' in lc:
        return 'CRANES', 'MONORAIL HOIST CRANE'
    if 'LIFE BOAT DAVIT' in ld or 'LIFEBOAT DAVIT' in ld or 'LBDS' in lc:
        return 'LIFE BOAT DAVITS', 'LIFE BOAT DAVIT'
    if 'RESCUE AND LIFERAFT DAVIT' in ld or 'RESLFDAV' in lc or 'RESCUE BOAT DAVIT' in ld:
        return 'LIFE BOAT DAVITS', 'RESCUE BOAT DAVIT'
    if 'E/R CRANE' in ld or 'ENGINE ROOM CRANE' in ld or 'ERCR' in lc:
        return 'CRANES', 'E/R OVERHEAD CRANE'
        
    # 8. Boats & Rafts
    if 'LIFEBOAT' in ld or 'LIFE BOAT' in ld or 'LIFBP' in lc:
        return 'LIFE BOATS', 'LIFE BOAT'
    if 'RESCUE BOAT' in ld or 'LRSCBO' in lc:
        return 'LIFE BOATS', 'RESCUE BOAT'
    if 'LIFERAFT' in ld or 'LIFE RAFT' in ld or 'LIFERA' in lc:
        return 'LIFE RAFTS', 'LIFE RAFT'
        
    # 9. Ladders
    if 'ACCOMMOD' in ld or 'ALRP' in lc or 'ALRS' in lc:
        return 'LIFTING APPLIANCES, DERRICKS, LADDERS', 'ACCOMMODATION LADDER'
    if 'PILOT' in ld or 'PISLAD' in lc:
        return 'LIFTING APPLIANCES, DERRICKS, LADDERS', 'PILOT LADDER'
        
    # 10. Cargo Hatch Cover
    if 'HATCH COVER' in ld or 'CAHACOV' in lc:
        return 'HOLDS', 'CARGO HATCH COVER'
        
    # 11. Fans & Ventilation
    if 'FAN' in ld or 'VENT' in ld:
        return 'HEATING, VENTILATION & AIR CONDITIONING', 'VENTILATION FAN'
        
    # 12. Galley & Laundry
    if 'GALLEY' in ld or 'LAUNDRY' in ld:
        return 'GALLEY', 'GALLEY & LAUNDRY EQUIPMENT'
        
    # 13. Window Wiper & CVS
    if 'CLEAR VIEW' in ld or 'CVS' in ld:
        return 'NAVIGATION EQUIPMENT', 'WINDOW WIPER & CLEAR VIEW SCREEN'
    if 'WINDOW WIPER' in ld or 'WIPER' in ld:
        return 'NAVIGATION EQUIPMENT', 'WINDOW WIPER & CLEAR VIEW SCREEN'
        
    # 14. Air Conditioning & Refrigeration
    if 'AIR CONDITIONING UNIT' in ld or 'AIR HANDLING UNIT' in ld or 'A/C PLANT' in ld:
        return 'HEATING, VENTILATION & AIR CONDITIONING', 'AIR CONDITIONING PLANT'
    if 'AIR CONDITION FOR ENGINE CONTROL ROOM' in ld or 'ECR' in ld and 'AIR' in ld:
        return 'AIR CONDITION', 'ECR PACKAGED AC'
    if 'REFRIGERATING' in ld or 'PROVISION REF' in ld or 'PROREF' in lc:
        return 'REFRIGERATION', 'PROVISION REFRIGERATING PLANT'
    if 'ROOM UNIT COOLER' in ld or 'COOLER' in ld and ('FISH' in ld or 'MEAT' in ld or 'VEGETABLES' in ld):
        return 'REFRIGERATION', 'COLD PROVISION CHAMBER'
        
    # 15. Elevator
    if 'ELEVATOR' in ld:
        return 'LIFTING APPLIANCES, DERRICKS, LADDERS', 'ELEVATOR'
        
    # 16. Fire Fighting
    if 'CO2' in ld and ('FIRE' in ld or 'EXT' in ld or 'SYSTEM' in ld):
        return 'FIRE FIGHTING', 'CO2 EXTINGUISHING SYSTEM'
    if 'FIRE ALARM' in ld or 'FIRE DETECTION' in ld or 'FIDES' in lc:
        return 'FIRE FIGHTING', 'FIRE ALARM & DETECTION SYSTEM'
    if 'LOCAL FIRE' in ld or 'FIRE FIGHTING' in ld or 'LOFFS' in lc:
        return 'FIRE FIGHTING', 'LOCAL FIRE EXTINGUISHING SYSTEM'
        
    # 17. Valves & Remote Control
    if 'VALVE REMOTE CONTROL' in ld or 'VRCS' in lc:
        return 'HYDRAULIC CONTROL UNIT', 'VALVE REMOTE CONTROL SYSTEM'
    if 'SHIPSIDE VALVES' in ld or 'SHIP SIDE' in ld:
        return 'SHIP SIDE VALVES', 'SEA CHEST & SHIP SIDE VALVE'
    if 'PRESSURE CONTROL VALVE' in ld or 'CONTROL VALVE' in ld:
        return 'VALVES', 'CONTROL VALVES'
    if 'EMERGENCY SHUT-OFF VALVE' in ld or 'SHUT-OFF' in ld:
        return 'VALVES', 'EMERGENCY SHUT-OFF VALVE'
        
    # 18. Gauging & Monitoring
    if 'LEVEL & DRAFT GAUGING' in ld or 'TANK LEVEL' in ld or 'ELEC. PNEUMATIC TYPE' in sd:
        return 'TANK LEVEL, GAUGING & MONITORING SYSTEM', 'ELEC. PNEUMATIC TYPE TANK LEVEL GAUGE'
    if 'LEVEL SWITCH' in ld or 'LEVEL SENSOR' in ld:
        return 'LEVEL INDICATORS', 'LEVEL SWITCH'
    if 'ALARM & MONITORING' in ld or 'AMS' in lc:
        return 'TANK LEVEL, GAUGING & MONITORING SYSTEM', 'ER ALARM & MONITORING SYS'
        
    # 19. Shafting & Stern Tube
    if 'STERN TUBE SEAL' in ld:
        return 'PROPELLER, THRUSTER, STERN TUBE, SHAFTING', 'STERN TUBE SEAL'
    if 'STERN TUBE BEARING' in ld:
        return 'PROPELLER, THRUSTER, STERN TUBE, SHAFTING', 'STERN TUBE BEARINGS'
    if 'INTERMEDIATE SHAFT' in ld or 'PROPELLER SHAFT' in ld or 'SHAFT' in ld and 'BEARING' in ld:
        return 'PROPELLER, THRUSTER, STERN TUBE, SHAFTING', 'INTERMEDIATE SHAFT BEARING'
    if 'PROPELLER' in ld and 'PUMP' not in ld:
        return 'PROPELLER, THRUSTER, STERN TUBE, SHAFTING', 'PROPELLER'
    if 'TOP BRACING' in ld:
        return 'MAIN ENGINE', 'MAIN ENGINE TOP BRACING'
        
    # 20. Engines & Generators
    if 'MAIN ENGINE' in ld or 'M/E' in ld and ('ENGINE' in ld or lc == 'ME'):
        return 'MAIN ENGINE', 'MAIN ENGINE'
    if 'DIESEL GENERATOR' in ld and ('ALTERNATOR' not in ld and 'FAN' not in ld and 'CHOCK' not in ld):
        return 'DIESEL GENERATOR', 'DIESEL GENERATOR ENGINE'
    if 'D/G ENGINE ALTERNATOR' in ld or 'MAIN DIESEL GENERATOR' in sd and 'ALTERNATOR' in ld:
        return 'ALTERNATOR', 'SYNCHRONOUS GENERATOR (D/G ALTERNATOR)'
    if 'EMERGENCY GENERATOR' in ld or 'EM\'CY DIESEL GENERATOR' in sd:
        return 'EMERGENCY GENERATOR', 'EM\'CY DIESEL GENERATOR SET'
    if 'EMERGENCY ALTERNATOR' in ld or 'EM\'CY GENERATOR SET ALTERNATOR' in sd:
        return 'EMERGENCY GENERATOR', 'EM\'CY DIESEL GENERATOR ALTERNATOR'
        
    # 21. Air Compressors & Reservoirs
    if 'AIR COMPR' in ld or 'AIR COMPRESSOR' in ld:
        if 'EM\'CY' in ld or 'EMERGENCY' in ld:
            return 'AIR COMPRESSOR', 'EMERGENCY AIR COMPRESSOR'
        if 'WORKING' in ld:
            return 'AIR COMPRESSOR', 'WORKING AIR COMPRESSOR'
        return 'AIR COMPRESSOR', 'MAIN AIR COMPRESSOR'
    if 'AIR RESERV' in ld or 'AIR RESERVOIR' in ld:
        return 'AIR RESERVOIR', 'AIR RESEVOIR'
    if 'AIR DRYER' in ld:
        return 'AIR DRYER', 'CONTROL AIR DRYER'
        
    # 22. Purifiers
    if 'PURIFIER' in ld and 'HEATER' not in ld and 'PUMP' not in ld:
        if 'L.F.O' in ld or 'H.F.O' in ld or 'FO' in ld:
            return 'PURIFIER', 'LFO PURIFIER'
        return 'PURIFIER', 'MAIN LO PURIFIER'
        
    # 23. Boilers, Coolers & Heaters
    if 'BOILER' in ld and 'COOLER' not in ld and 'PUMP' not in ld:
        return 'BOILER', 'COMPOSITE BOILER'
    if 'HEAT EXCHANGER' in ld or 'COOLER' in ld:
        return 'COOLERS', 'COOLER'
    if 'HEATER' in ld or 'PREHEATER' in ld:
        return 'F.W. GENERATOR, HEAT EXCHANGER', 'HEATER'
    if 'FRESH WATER GENERATOR' in ld or 'F.W. GENERATOR' in ld:
        return 'F.W. GENERATOR, HEAT EXCHANGER', 'F.W. GENERATOR'
        
    # 24. Workshop
    if 'LATHE' in ld:
        return 'WORKSHOP MACHINERY, TOOLS, PORTABLE INSTRUMENTS', 'LATHE MACHINE'
    if 'GRINDER' in ld or 'GRINDING' in ld:
        return 'WORKSHOP MACHINERY, TOOLS, PORTABLE INSTRUMENTS', 'DRILLING & GRINDING MACHINE'
    if 'DRILLING' in ld:
        return 'WORKSHOP MACHINERY, TOOLS, PORTABLE INSTRUMENTS', 'DRILLING & GRINDING MACHINE'
    if 'GAS WELDER' in ld:
        return 'WORKSHOP MACHINERY, TOOLS, PORTABLE INSTRUMENTS', 'OXYGEN ACETYLENE GAS WELDING'
    if 'ELECTRIC ARC WELDER' in ld or 'ELECTRIC WELDER' in ld:
        return 'ELECTRICAL SYSTEM', 'ELECTRIC WELDING MACHINE'
        
    # 25. Domestic & Sanitary
    if 'SEWAGE' in ld and 'PUMP' not in ld:
        return 'SEWAGE SYSTEM', 'SEWAGE TREATMENT PLANT'
    if 'F.W. SUPPLY' in ld or 'HYDROPHORE' in ld:
        return 'HYDROPHORE TANK', 'FW HYDROPHORE SUPPLY UNIT'
    if 'TOILET' in ld:
        return 'VACUUM TOILET SYSTEM & SANITARY EQUIPMENT', 'VACUUM TOILET'
        
    # 26. Environmental & Fuel Conditioning
    if 'ANTI-FOULING' in ld or 'MGPS' in ld:
        return 'M.G.P.S, I.C.C.P., ANODES, ANTI-FOULING', 'MGPS'
    if 'BALLAST WATER' in ld or 'BWMS' in ld:
        return 'BALLAST WATER TREATMENT SYSTEM', 'BALLAST WATER TREATMENT PLANT'
    if 'BILGE SEPARATOR' in ld or 'OILY BILGE' in ld or '5 PPM' in ld:
        return 'OIL SPILL EQUIPMENT', 'OILY BILGE SEPARATOR'
    if 'INCINERATOR' in ld:
        return 'INCINERATORS', 'INCINERATOR'
    if 'F.O.SUPPLY UNIT' in ld or 'F.O SUPPLY' in ld:
        return 'MAIN ENGINE', 'F.O SUPPLY UNIT FOR ME & GE'
    if 'EXHAUST GAS CLEANING' in ld or 'SCR' in ld:
        return 'MAIN ENGINE', 'HP SCR'
    if 'AUTO L.O. FILTER' in ld or 'FILTER' in ld:
        return 'PIPING, VALVES, FILTERS, EJECTOR, FITTINGS', 'ME LO AUTO-BACK FLUSHINGFINE FILTER'
        
    # 27. Pumps
    if 'PUMP' in ld:
        if 'ANTI-HEELING' in ld:
            return 'PUMP', 'ANTI-HEELING PUMP'
        if 'EMERGENCY FIRE' in ld:
            return 'PUMP', 'EMERGENCY FIRE PUMP'
        return 'PUMP', 'ALL PUMPS'
        
    # 28. Switchboards, Transformers & Electrical
    if 'SWITCHBOARD' in ld or 'SWITCH BOARD' in ld:
        if 'EMERGENCY' in ld or 'EM\'CY' in ld:
            return 'SWITCHBOARDS', 'EM\'CY SWITCH BOARD'
        return 'SWITCHBOARDS', 'MAIN SWITCH BOARD'
    if 'DISTRIBUTION BOARD' in ld or 'DNBD' in lc:
        return 'ELECTRICAL SYSTEM', 'DISTRIBUTION BOARD'
    if 'STARTER' in ld or 'PANEL' in ld:
        return 'ELECTRICAL SYSTEM', 'GROUP STARTER & FEEDER PANEL'
    if 'STOP SWITCH' in ld:
        return 'ELECTRICAL SYSTEM', 'EMERGENCY STOP SWITCH BOX(ELECTRICAL)'
    if 'TRANSFORMER' in ld:
        return 'ELECTRICAL SYSTEM', 'POWER & LIGHTING TRANSFORMER'
    if 'REEFER' in ld:
        if 'MONITORING' in ld:
            return 'CARGO CONTAINER SYSTEM', 'REEFER MONITORING SYSTEM'
        return 'CARGO CONTAINER SYSTEM', 'REEFER CONTAINER RECEPTACLE'
    if 'LIGHT' in ld or 'LAMP' in ld:
        return 'ELECTRICAL SYSTEM', 'LIGHTING FIXTURE'
    if 'RECEPTACLE' in ld or 'PLUG' in ld:
        return 'ELECTRICAL SYSTEM', 'ELEC. EQUIP. & CABLE WAY(ELECTRICAL)'
    if 'SIGNAL LIGHT COLUMN' in ld or 'LIGHT SIGNAL' in ld:
        return 'ELECTRONIC EQUIPMENT', 'LIGHT SIGNAL COLUMN'
    if 'BATTERY CHARGER' in ld:
        return 'ELECTRONIC EQUIPMENT', 'BATTERY CHARGER & DIST. BOARD'
    if 'BATTERY' in ld:
        return 'BATTERIES', 'BATTERIES'
    if 'GAUGE BOARD' in ld:
        return 'ELECTRICAL SYSTEM', 'WHEEL HOUSE GAUGE BOARD'
        
    # 29. Instrumentation & Automation
    if 'HORSEPOWER METER' in ld or 'HORSE POWER' in ld:
        return 'MEASUREMENT INSTRUMENTS', 'ME HORSE POWER METER'
    if 'WING CONSOLE' in ld:
        return 'E/R CONSOLE', 'BRIDGE WING CONTROL CONSOLE'
    if 'BRIDGE CONTROL CONSOLE' in ld or 'ECC' in lc:
        return 'E/R CONSOLE', 'ENGINE CONTROL CONSOLE'
    if 'SHIP\'S COMPUTER' in ld or 'PERFORMANCE SYSTEM' in sd or 'ISS' in sd:
        return 'MEASUREMENT INSTRUMENTS', 'INTEGRATED MONITORING AND CONTROL SYSTEM'
    if 'CCTV' in ld:
        return 'ELECTRONIC EQUIPMENT', 'CCTV'
    if 'EARTHING' in ld or 'GROUNDING' in ld:
        return 'ELECTRONIC EQUIPMENT', 'SHAFT EARTHING DEVICE'
        
    # 30. Navigation & Communication
    if 'RADAR' in ld:
        if 'TRANSPONDER' in ld or 'SART' in ld:
            return 'NAVIGATION EQUIPMENT', 'SART'
        return 'RDR', 'RADAR EQUIPMENT'
    if 'BRIDGE INTERFACE' in ld:
        return 'NAVIGATION EQUIPMENT', 'CONNING SYSTEM'
    if 'GPS' in ld or 'DGPS' in ld:
        return 'NAVIGATION EQUIPMENT', 'GPS'
    if 'SATELLITE LOG' in ld or 'SPEED LOG' in ld:
        return 'NAVIGATION EQUIPMENT', 'SPEED LOG'
    if 'MF/HF' in ld or 'GMDSS' in ld:
        return 'COMMUNICATIONS', 'MF HF TRANSRECEIVER'
    if 'AIS' in ld:
        return 'NAVIGATION EQUIPMENT', 'AIS'
    if 'CONNING' in ld:
        return 'NAVIGATION EQUIPMENT', 'CONNING SYSTEM'
    if 'INM C' in ld or 'INMARSAT' in ld:
        return 'NAVIGATION EQUIPMENT', 'INMARSAT C'
    if 'NAVTEX' in ld:
        return 'NAVIGATION EQUIPMENT', 'NAVTEX'
    if 'EPIRB' in ld:
        return 'NAVIGATION EQUIPMENT', 'EPIRB'
    if 'TWO-WAY VHF' in ld or 'PORTABLE VHF' in sd:
        return 'COMMUNICATIONS', '2 WAY VHF'
    if 'VHF' in ld:
        return 'COMMUNICATIONS', 'VHF'
    if 'VOYAGE DATA' in ld or 'VDR' in ld:
        return 'NAVIGATION EQUIPMENT', 'VDR'
    if 'WEATHER FAX' in ld or 'FAX' in ld:
        return 'NAVIGATION EQUIPMENT', 'WEATHER FACSIMILE'
    if 'GYRO' in ld:
        return 'NAVIGATION EQUIPMENT', 'GYRO COMPASS'
    if 'MAGNETIC COMPASS' in ld:
        return 'NAVIGATION EQUIPMENT', 'MAGNETIC COMPASS'
    if 'ECDIS' in ld:
        return 'NAVIGATION EQUIPMENT', 'ECDIS'
    if 'ECHO SOUNDER' in ld:
        return 'NAVIGATION EQUIPMENT', 'ECHO SOUNDER'
    if 'AUTOPILOT' in ld or 'AUTO PILOT' in ld:
        return 'NAVIGATION EQUIPMENT', 'AUTO PILOT'
    if 'BNWAS' in ld or 'BRIDGE NAVIGATION WATCH' in ld:
        return 'NAVIGATION EQUIPMENT', 'BRIDGE WATCH ALARM SYSTEM'
    if 'RUDDER ANGLE' in ld or 'RUDANGIN' in lc:
        return 'RUDDER', 'RUDDER ANGLE INDICATOR'
    if 'ANEMOMETER' in ld or 'AMAS' in lc:
        return 'NAVIGATION EQUIPMENT', 'ANEMOMETER & ANEMOSCOPE'
    if 'WHISTLE' in ld:
        return 'WHISTLES', 'WHISTLE'
    if 'PUBLIC ADDRESS' in ld or 'PA SYSTEM' in ld:
        return 'COMMUNICATIONS', 'COMMUNICATION SYSTEM'
    if 'TELEPHONE' in ld:
        return 'COMMUNICATIONS', 'COMMUNICATION SYSTEM'
    if 'NETWORK' in ld or 'AERIAL' in ld or 'IRIDIUM' in ld or 'V-SAT' in ld or 'VSAT' in ld or 'UHF' in ld:
        return 'COMMUNICATIONS', 'COMMUNICATION SYSTEM'
    if 'CLOCK' in ld:
        return 'COMMUNICATIONS', 'COMMUNICATION SYSTEM'
    if 'HOSPITAL CALLING' in ld or 'REF. CHAMBER ALARM' in ld:
        return 'FIRE FIGHTING', 'HOSPITAL & REF CHAMBER ALARM'
        
    return None, None

results = []
missing = []
for r in range(2, ws_main.max_row + 1):
    loc_code = str(ws_main.cell(r, 2).value or '')
    loc_desc = str(ws_main.cell(r, 3).value or '')
    sys_desc = str(ws_main.cell(r, 5).value or '')
    maker = str(ws_main.cell(r, 7).value or '')
    model = str(ws_main.cell(r, 16).value or '')
    
    fn, mach = get_function_for_row(r, loc_code, loc_desc, sys_desc, maker, model)
    if fn:
        results.append((r, loc_code, loc_desc, fn, mach))
    else:
        missing.append((r, loc_code, loc_desc, sys_desc))

print(f'Successfully mapped: {len(results)} / {ws_main.max_row - 1}')
print(f'Missing: {len(missing)}')
if missing:
    print('Missing rows:')
    for m in missing:
        print(f'Row {m[0]}: [{m[1]}] {m[2]} | {m[3]}')

with open('01_VL_EXTRACTION/mapped_results.txt', 'w', encoding='utf-8') as f:
    for r, lc, ld, fn, mach in results:
        f.write(f'Row {r:3d} | [{lc:8s}] {ld:42s} -> Function: {fn:35s} (Mach: {mach})\n')

print('Saved mapped_results.txt')
