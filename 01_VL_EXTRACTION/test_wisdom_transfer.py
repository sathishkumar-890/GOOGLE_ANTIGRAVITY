import json
import openpyxl

with open('01_VL_EXTRACTION/reward_fun_items.json', 'r', encoding='utf-8') as f:
    fun_items = json.load(f)

with open('01_VL_EXTRACTION/reward_main_vdc_rows.json', 'r', encoding='utf-8') as f:
    main_rows = json.load(f)

# Build map of machinery name -> list of items
fun_mach_to_items = {}
for it in fun_items:
    m = it['machinery'].strip().upper()
    if m not in fun_mach_to_items:
        fun_mach_to_items[m] = []
    fun_mach_to_items[m].append(it)

# Also collect all unique functions in Reward FUN
valid_funcs = set(it['function'].strip() for it in fun_items)

print(f"Total valid functions in REWARD FUN: {len(valid_funcs)}")

# Let's inspect Wisdom's mapped results and see how they translate to Reward
wb_w = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC.xlsx', data_only=True)
ws_w = wb_w.active
wisdom_map = {} # loc_code -> function
for r in range(2, ws_w.max_row + 1):
    lc = ws_w.cell(r, 2).value
    fn = ws_w.cell(r, 6).value
    if lc and fn:
        wisdom_map[str(lc).strip().upper()] = str(fn).strip()

print(f"Wisdom mapped locations: {len(wisdom_map)}")

# Let's see how many of Wisdom's functions directly exist in Reward's valid_funcs
w_in_r = 0
w_not_in_r = set()
for r in main_rows:
    lc = str(r['loc_code'] or '').strip().upper()
    w_fn = wisdom_map.get(lc)
    if w_fn:
        if w_fn in valid_funcs:
            w_in_r += 1
        else:
            w_not_in_r.add(w_fn)

print(f"Matches where Wisdom function is directly valid in Reward: {w_in_r} / {len(main_rows)}")
print(f"Wisdom functions that are NOT in Reward: {w_not_in_r}")
