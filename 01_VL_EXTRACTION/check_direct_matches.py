import openpyxl

wb_fun = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws_fun = wb_fun['Allocation Sheet']

fun_dict = {}
for r in range(12, ws_fun.max_row + 1):
    func = ws_fun.cell(r, 9).value
    mach = ws_fun.cell(r, 10).value
    hyper = ws_fun.cell(r, 3).value
    if mach and mach != 'Folder' and func and func != 'Folder':
        m_str = str(mach).strip().upper()
        f_str = str(func).strip().upper()
        # if duplicate, store first or keep both
        if m_str not in fun_dict:
            fun_dict[m_str] = (str(func).strip(), str(mach).strip())

wb_main = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx', data_only=True)
ws_main = wb_main['MAIN VDC_REWARD']

print(f"Total rows in MAIN VDC_REWARD: {ws_main.max_row}")

rows_data = []
for r in range(2, ws_main.max_row + 1):
    loc_code = str(ws_main.cell(r, 2).value or '').strip()
    loc_desc = str(ws_main.cell(r, 3).value or '').strip()
    sys_code = str(ws_main.cell(r, 4).value or '').strip()
    sys_desc = str(ws_main.cell(r, 5).value or '').strip()
    maker = str(ws_main.cell(r, 7).value or '').strip()
    model = str(ws_main.cell(r, 16).value or '').strip()
    rows_data.append({
        'row': r,
        'loc_code': loc_code,
        'loc_desc': loc_desc,
        'sys_code': sys_code,
        'sys_desc': sys_desc,
        'maker': maker,
        'model': model
    })

print(f"Loaded {len(rows_data)} equipment rows.")

# Let's see how many direct matches we get if we match loc_desc against machinery names in fun_dict
direct_matches = 0
unmatched = []
for rd in rows_data:
    ld = rd['loc_desc'].upper()
    if ld in fun_dict:
        direct_matches += 1
    else:
        unmatched.append(rd)

print(f"Exact direct matches on loc_desc: {direct_matches} / {len(rows_data)}")
print(f"Unmatched on direct loc_desc: {len(unmatched)}")
