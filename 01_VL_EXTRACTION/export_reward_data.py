import openpyxl
import json

wb = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws = wb['Allocation Sheet']

items = []
for r in range(12, ws.max_row + 1):
    func = ws.cell(r, 9).value
    mach = ws.cell(r, 10).value
    hyper = ws.cell(r, 3).value
    parent = ws.cell(r, 2).value
    item_type = ws.cell(r, 5).value
    
    if mach and mach != 'Folder' and func and func != 'Folder':
        items.append({
            'row': r,
            'function': str(func).strip(),
            'machinery': str(mach).strip(),
            'hyperlink': str(hyper).strip() if hyper else '',
            'parent': str(parent).strip() if parent else '',
            'item_type': str(item_type).strip() if item_type else ''
        })

with open('01_VL_EXTRACTION/reward_fun_items.json', 'w', encoding='utf-8') as f:
    json.dump(items, f, indent=2)

print(f"Saved {len(items)} items to reward_fun_items.json")

# Also let's extract all rows from MAIN VDC_REWARD
wb_main = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx', data_only=True)
ws_main = wb_main['MAIN VDC_REWARD']

main_rows = []
for r in range(2, ws_main.max_row + 1):
    main_rows.append({
        'row': r,
        'vessel_code': ws_main.cell(r, 1).value,
        'loc_code': ws_main.cell(r, 2).value,
        'loc_desc': ws_main.cell(r, 3).value,
        'sys_code': ws_main.cell(r, 4).value,
        'sys_desc': ws_main.cell(r, 5).value,
        'func_desc': ws_main.cell(r, 6).value,
        'maker': ws_main.cell(r, 7).value,
        'critical': ws_main.cell(r, 8).value,
        'serial': ws_main.cell(r, 9).value,
        'particulars': ws_main.cell(r, 10).value,
        'manual_name': ws_main.cell(r, 14).value,
        'page_no': ws_main.cell(r, 15).value,
        'model': ws_main.cell(r, 16).value
    })

with open('01_VL_EXTRACTION/reward_main_vdc_rows.json', 'w', encoding='utf-8') as f:
    json.dump(main_rows, f, indent=2)

print(f"Saved {len(main_rows)} rows to reward_main_vdc_rows.json")
