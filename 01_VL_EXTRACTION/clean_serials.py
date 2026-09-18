import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import shutil
import os

EXCEL_PATH = '01_VL_EXTRACTION/VL_IRENES REWARD.xlsx'
ROOT_EXCEL_PATH = 'VL_IRENES REWARD.xlsx'
BACKUP_PATH = '01_VL_EXTRACTION/VL_IRENES REWARD_BACKUP_BEFORE_CLEAN.xlsx'

shutil.copy2(EXCEL_PATH, BACKUP_PATH)
print(f'Backup saved to {BACKUP_PATH}')

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
ws = wb.active

green_fill = PatternFill(start_color='FF92D050', end_color='FF92D050', fill_type='solid')
no_fill = PatternFill(fill_type=None)

thin_border = Border(
    left=Side(style='thin', color='FF000000'),
    right=Side(style='thin', color='FF000000'),
    top=Side(style='thin', color='FF000000'),
    bottom=Side(style='thin', color='FF000000')
)
data_font = Font(name='Calibri', size=11)
left_align = Alignment(horizontal='left', vertical='center')
center_align = Alignment(horizontal='center', vertical='center')

dg_engines = {
    92: ('KBA008085-1', 'VOL 2.pdf', '643'),
    93: ('KBA008085-2', 'VOL 2.pdf', '644'),
    94: ('KBA008085-3', 'VOL 2.pdf', '645'),
}

cleared_count = 0
kept_count = 0

for r in range(2, ws.max_row + 1):
    c8 = ws.cell(row=r, column=8)
    c9 = ws.cell(row=r, column=9)
    c10 = ws.cell(row=r, column=10)

    # Standardize borders and font
    c8.border = thin_border
    c9.border = thin_border
    c10.border = thin_border
    c8.font = data_font
    c9.font = data_font
    c10.font = data_font

    if r in dg_engines:
        s_no, pdf_name, page_no = dg_engines[r]
        c8.value = s_no
        c8.fill = green_fill
        c8.alignment = center_align
        c9.value = pdf_name
        c9.alignment = left_align
        c10.value = page_no
        c10.alignment = center_align
        kept_count += 1
    else:
        # Clear drawing numbers and assumptions
        if c8.value is not None or c9.value is not None or c10.value is not None:
            cleared_count += 1
        c8.value = None
        c8.fill = no_fill
        c8.alignment = None
        c9.value = None
        c9.alignment = left_align
        c10.value = None
        c10.alignment = center_align

print(f'Kept verified engine rows: {kept_count}')
print(f'Cleared non-engine rows: {cleared_count}')

wb.save(EXCEL_PATH)
print(f'Saved updated workbook to {EXCEL_PATH}')

shutil.copy2(EXCEL_PATH, ROOT_EXCEL_PATH)
print(f'Synchronized root replica to {ROOT_EXCEL_PATH}')
