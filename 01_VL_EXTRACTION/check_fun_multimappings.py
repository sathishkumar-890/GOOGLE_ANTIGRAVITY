import openpyxl

wb_fun = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws_fun = wb_fun['Allocation Sheet']

entries = []
for r in range(12, ws_fun.max_row + 1):
    func = ws_fun.cell(r, 9).value
    mach = ws_fun.cell(r, 10).value
    hyper = ws_fun.cell(r, 3).value
    parent = ws_fun.cell(r, 2).value
    if mach and mach != 'Folder' and func and func != 'Folder':
        entries.append({
            'row': r,
            'function': str(func).strip(),
            'machinery': str(mach).strip(),
            'hyperlink': str(hyper).strip() if hyper else '',
            'parent': str(parent).strip() if parent else ''
        })

print(f'Total items: {len(entries)}')

# Check for duplicates or multiple functions per machinery
mach_to_funcs = {}
for e in entries:
    m = e['machinery']
    f = e['function']
    if m not in mach_to_funcs:
        mach_to_funcs[m] = set()
    mach_to_funcs[m].add(f)

multi_funcs = {m: list(fs) for m, fs in mach_to_funcs.items() if len(fs) > 1}
print(f'Machineries with multiple functions: {len(multi_funcs)}')
for m, fs in sorted(multi_funcs.items()):
    print(f'  {m}: {fs}')
