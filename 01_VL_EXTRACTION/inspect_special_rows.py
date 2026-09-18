import json
import openpyxl

with open('01_VL_EXTRACTION/reward_fun_items.json', 'r', encoding='utf-8') as f:
    fun_items = json.load(f)

with open('01_VL_EXTRACTION/reward_main_vdc_rows.json', 'r', encoding='utf-8') as f:
    main_rows = json.load(f)

wb_w = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC.xlsx', data_only=True)
ws_w = wb_w.active
wisdom_map = {}
for r in range(2, ws_w.max_row + 1):
    lc = ws_w.cell(r, 2).value
    fn = ws_w.cell(r, 6).value
    if lc and fn:
        wisdom_map[str(lc).strip().upper()] = str(fn).strip()

valid_funcs = set(it['function'].strip() for it in fun_items)

special_rows = []
for r in main_rows:
    lc = str(r['loc_code'] or '').strip().upper()
    w_fn = wisdom_map.get(lc)
    if not w_fn or w_fn not in valid_funcs:
        special_rows.append((r, w_fn))

print(f"Total special rows: {len(special_rows)}")
for r, w_fn in special_rows:
    print(f"Row {r['row']}: [{r['loc_code']}] {r['loc_desc']} | Wisdom_fn='{w_fn}'")
