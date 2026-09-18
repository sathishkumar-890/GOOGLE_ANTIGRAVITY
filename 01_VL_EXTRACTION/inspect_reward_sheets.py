import openpyxl

wb = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx', data_only=True)
ws_main = wb['MAIN VDC_REWARD']
ws_r2 = wb['Reward (2)']

print('MAIN VDC_REWARD headers:')
for c in range(1, ws_main.max_column + 1):
    print(f'  Col {c} ({openpyxl.utils.get_column_letter(c)}): {ws_main.cell(1, c).value}')

print('\nReward (2) headers:')
for c in range(1, ws_r2.max_column + 1):
    print(f'  Col {c} ({openpyxl.utils.get_column_letter(c)}): {ws_r2.cell(1, c).value}')

print('\nChecking first 5 rows of MAIN VDC_REWARD:')
for r in range(2, 7):
    row_vals = [ws_main.cell(r, c).value for c in range(1, 17)]
    print(f'Row {r}:', row_vals)

print('\nChecking first 5 rows of Reward (2):')
for r in range(2, 7):
    row_vals = [ws_r2.cell(r, c).value for c in range(1, 11)]
    print(f'Row {r}:', row_vals)
