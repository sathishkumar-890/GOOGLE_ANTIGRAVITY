import openpyxl

wb_fun = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws_fun = wb_fun['Allocation Sheet']

fun_pairs = []
unique_funcs = set()
unique_machs = set()
for r in range(12, ws_fun.max_row + 1):
    func = ws_fun.cell(r, 9).value
    mach = ws_fun.cell(r, 10).value
    hyper = ws_fun.cell(r, 3).value
    item_type = ws_fun.cell(r, 5).value
    if mach and mach != 'Folder' and func and func != 'Folder':
        mach_str = str(mach).strip()
        func_str = str(func).strip()
        fun_pairs.append((mach_str, func_str, str(hyper).strip() if hyper else '', r))
        unique_funcs.add(func_str)
        unique_machs.add(mach_str)

print(f'Total valid Item rows in FUN: {len(fun_pairs)}')
print(f'Total unique Functions in FUN: {len(unique_funcs)}')
print(f'Total unique Machineries in FUN: {len(unique_machs)}')
print('\nUnique functions in FUN:')
for f in sorted(unique_funcs):
    print(f'  - {f}')

wb_main = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx', data_only=True)
print('\nActive sheet in MAIN VDC:', wb_main.active.title)

for sname in ['MAIN VDC_REWARD', 'Reward (2)', 'Reward']:
    if sname in wb_main.sheetnames:
        ws = wb_main[sname]
        col_f_name = ws.cell(1, 6).value
        non_empty = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 6).value not in [None, ''])
        empty = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 6).value in [None, ''])
        print(f'Sheet {sname}: max_row={ws.max_row}, max_col={ws.max_column}, Col 6 Header="{col_f_name}", Non-empty={non_empty}, Empty={empty}')
