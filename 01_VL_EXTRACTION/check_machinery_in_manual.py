import openpyxl

wb = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx', data_only=True)
ws_rew = wb['Reward']
ws_r2 = wb['Reward (2)']
ws_main = wb['MAIN VDC_REWARD']

print("Row 2 comparison:")
print("Reward Col 3 (Machinery name in Manual):", ws_rew.cell(2, 3).value)
print("Reward (2) Col 18 (Machinery name in Manual):", ws_r2.cell(2, 18).value)

wb_fun = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws_fun = wb_fun['Allocation Sheet']
fun_dict = {}
for r in range(12, ws_fun.max_row + 1):
    func = ws_fun.cell(r, 9).value
    mach = ws_fun.cell(r, 10).value
    if mach and mach != 'Folder' and func and func != 'Folder':
        m_str = str(mach).strip().upper()
        if m_str not in fun_dict:
            fun_dict[m_str] = str(func).strip()

matches_col3 = 0
for r in range(2, ws_rew.max_row + 1):
    m = str(ws_rew.cell(r, 3).value or '').strip().upper()
    if m in fun_dict:
        matches_col3 += 1

print(f"Matches using 'Machinery name in Manual': {matches_col3} / {ws_rew.max_row - 1}")

# Sample of Col 3 vs fun_dict
for r in range(2, 20):
    m = str(ws_rew.cell(r, 3).value or '').strip()
    ld = str(ws_rew.cell(r, 2).value or '').strip()
    print(f"Row {r}: LocDesc='{ld}' | MachManual='{m}' | in_fun={m.upper() in fun_dict}")
