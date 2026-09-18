import json
import openpyxl

with open('01_VL_EXTRACTION/reward_main_vdc_rows.json', 'r', encoding='utf-8') as f:
    main_rows = json.load(f)

# Import mapping function
import sys
sys.path.append('01_VL_EXTRACTION')
from test_full_reward_mapping import map_reward_row

wb_w = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC.xlsx', data_only=True)
ws_w = wb_w.active
wisdom_map = {}
for r in range(2, ws_w.max_row + 1):
    lc = ws_w.cell(r, 2).value
    fn = ws_w.cell(r, 6).value
    if lc and fn:
        wisdom_map[str(lc).strip().upper()] = str(fn).strip()

diffs = []
all_results = []
for r in main_rows:
    fn, mach = map_reward_row(r)
    lc = str(r['loc_code'] or '').strip().upper()
    w_fn = wisdom_map.get(lc, 'N/A')
    all_results.append((r['row'], lc, r['loc_desc'], fn, mach, w_fn))
    if fn != w_fn:
        diffs.append((r['row'], lc, r['loc_desc'], fn, mach, w_fn))

print(f"Total rows: {len(all_results)}")
print(f"Rows where Reward Function differs from Wisdom Function: {len(diffs)}")
print("\nDetailed list of differences:")
for d in diffs:
    print(f"Row {d[0]:3d} | [{d[1]:8s}] {d[2]:35s} -> Reward: '{d[3]}' vs Wisdom: '{d[5]}'")

with open('01_VL_EXTRACTION/reward_mapping_full_audit.txt', 'w', encoding='utf-8') as f:
    for row_num, lc, ld, fn, mach, w_fn in all_results:
        f.write(f"Row {row_num:3d} | [{lc:8s}] {ld:42s} | FUNC: {fn:40s} | MACH: {mach:35s} | WISDOM_FUNC: {w_fn}\n")

print("\nSaved full audit to 01_VL_EXTRACTION/reward_mapping_full_audit.txt")
