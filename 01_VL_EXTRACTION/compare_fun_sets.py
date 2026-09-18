import openpyxl

wb_w = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES WISDOM_FUN.xlsx', data_only=True)
ws_w = wb_w.active
wisdom_funcs = set()
for r in range(7, ws_w.max_row + 1):
    f = ws_w.cell(r, 13).value
    if f and f != 'Folder':
        wisdom_funcs.add(str(f).strip())

wb_r = openpyxl.load_workbook('01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx', data_only=True)
ws_r = wb_r['Allocation Sheet']
reward_funcs = set()
for r in range(12, ws_r.max_row + 1):
    f = ws_r.cell(r, 9).value
    if f and f != 'Folder':
        reward_funcs.add(str(f).strip())

print(f"Wisdom unique functions: {len(wisdom_funcs)}")
print(f"Reward unique functions: {len(reward_funcs)}")

print("\nFunctions in Wisdom but NOT in Reward:")
for f in sorted(wisdom_funcs - reward_funcs):
    print(f"  - {f}")

print("\nFunctions in Reward but NOT in Wisdom:")
for f in sorted(reward_funcs - wisdom_funcs):
    print(f"  - {f}")
