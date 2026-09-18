import openpyxl

wb_r = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx', data_only=True)
ws_r_main = wb_r['MAIN VDC_REWARD']

wb_w = openpyxl.load_workbook('01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC.xlsx', data_only=True)
ws_w_main = wb_w.active

print(f"MAIN VDC_REWARD max_row: {ws_r_main.max_row}")
print(f"MAIN VDC_wisdom max_row: {ws_w_main.max_row}")

# Check row by row Location Code and Location Description
diffs = []
for r in range(2, max(ws_r_main.max_row, ws_w_main.max_row) + 1):
    lc_r = ws_r_main.cell(r, 2).value if r <= ws_r_main.max_row else None
    ld_r = ws_r_main.cell(r, 3).value if r <= ws_r_main.max_row else None
    
    lc_w = ws_w_main.cell(r, 2).value if r <= ws_w_main.max_row else None
    ld_w = ws_w_main.cell(r, 3).value if r <= ws_w_main.max_row else None
    
    if lc_r != lc_w or ld_r != ld_w:
        diffs.append((r, lc_r, ld_r, lc_w, ld_w))

print(f"Differences in (Location Code, Location Description) between Wisdom and Reward: {len(diffs)}")
for d in diffs[:10]:
    print(f"Row {d[0]}: Reward=({d[1]}, {d[2]}) vs Wisdom=({d[3]}, {d[4]})")
