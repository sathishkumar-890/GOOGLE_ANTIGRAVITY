import openpyxl
from openpyxl.styles import Font, Alignment
import shutil
import os

EXCEL_PATH = '01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC.xlsx'
BACKUP_PATH = '01_VL_EXTRACTION/VL_EXTRACTION_WISDOM-MAINVDC_BACKUP.xlsx'
FUN_PATH = '01_VL_EXTRACTION/C.V. IRENES WISDOM_FUN.xlsx'

shutil.copy2(EXCEL_PATH, BACKUP_PATH)
print(f'Backup saved to {BACKUP_PATH}')

wb_fun = openpyxl.load_workbook(FUN_PATH, data_only=True)
ws_fun = wb_fun.active

# Valid functions in FUN
valid_funcs = set()
for r in range(7, ws_fun.max_row + 1):
    f = ws_fun.cell(r, 13).value
    if f and f != 'Folder':
        valid_funcs.add(str(f).strip())

wb_main = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
ws_main = wb_main.active

# Header check
assert ws_main.cell(1, 6).value == 'Function Description', f"Expected Col 6 'Function Description', got {ws_main.cell(1, 6).value}"

# Load mapping logic from build_function_mapping
import sys
sys.path.append('01_VL_EXTRACTION')
from build_function_mapping import get_function_for_row

updated_count = 0
data_font = Font(name='Calibri', size=11)
align = Alignment(horizontal='left', vertical='center')

for r in range(2, ws_main.max_row + 1):
    loc_code = str(ws_main.cell(r, 2).value or '')
    loc_desc = str(ws_main.cell(r, 3).value or '')
    sys_desc = str(ws_main.cell(r, 5).value or '')
    maker = str(ws_main.cell(r, 7).value or '')
    model = str(ws_main.cell(r, 16).value or '')
    
    func, mach = get_function_for_row(r, loc_code, loc_desc, sys_desc, maker, model)
    assert func is not None, f'Row {r} could not be mapped: {loc_code} | {loc_desc}'
    assert func in valid_funcs, f'Row {r} function {func} not in FUN valid functions!'
    
    cell_f = ws_main.cell(row=r, column=6)
    cell_f.value = func
    cell_f.font = data_font
    cell_f.alignment = align
    updated_count += 1

print(f'Successfully updated {updated_count} rows in Function Description (Col F)')

# Set column width
ws_main.column_dimensions['F'].width = 38

wb_main.save(EXCEL_PATH)
print(f'Saved updated workbook to {EXCEL_PATH}')
