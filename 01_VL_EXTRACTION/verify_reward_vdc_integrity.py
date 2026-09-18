import openpyxl

BACKUP_PATH = '01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC_BACKUP.xlsx'
UPDATED_PATH = '01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx'
FUN_PATH = '01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx'

# Load FUN valid functions
wb_fun = openpyxl.load_workbook(FUN_PATH, data_only=True)
ws_fun = wb_fun['Allocation Sheet']
valid_funcs = set()
for r in range(12, ws_fun.max_row + 1):
    f = ws_fun.cell(r, 9).value
    if f and f != 'Folder':
        valid_funcs.add(str(f).strip())

print(f"Total valid functions in FUN: {len(valid_funcs)}")

wb_bk = openpyxl.load_workbook(BACKUP_PATH, data_only=False)
wb_up = openpyxl.load_workbook(UPDATED_PATH, data_only=False)

# 1. Verify MAIN VDC_REWARD
ws_bk_main = wb_bk['MAIN VDC_REWARD']
ws_up_main = wb_up['MAIN VDC_REWARD']

assert ws_up_main.max_row == 279, f"Expected 279 rows, got {ws_up_main.max_row}"
assert ws_up_main.max_column == 16, f"Expected 16 columns, got {ws_up_main.max_column}"

empty_count = 0
invalid_count = 0
for r in range(2, 280):
    val = ws_up_main.cell(r, 6).value
    if val is None or str(val).strip() == '':
        empty_count += 1
    elif str(val).strip() not in valid_funcs:
        invalid_count += 1
        print(f"Invalid function at row {r}: {val}")

assert empty_count == 0, f"Found {empty_count} empty cells in Col F!"
assert invalid_count == 0, f"Found {invalid_count} invalid functions in Col F!"
print("MAIN VDC_REWARD Col F: 278 / 278 rows populated with 100% valid functions.")

diffs = 0
for r in range(1, 280):
    for c in range(1, 17):
        if c == 6:
            continue
        v_bk = ws_bk_main.cell(r, c).value
        v_up = ws_up_main.cell(r, c).value
        if v_bk != v_up:
            diffs += 1
            print(f"Mismatch at ({r}, {c}): backup='{v_bk}' vs updated='{v_up}'")

assert diffs == 0, f"Found {diffs} unexpected differences in other columns!"
print("MAIN VDC_REWARD other 15 columns: 100% identical to backup.")

# 2. Verify Reward (2)
ws_bk_r2 = wb_bk['Reward (2)']
ws_up_r2 = wb_up['Reward (2)']

assert ws_up_r2.max_row == 279, f"Expected 279 rows, got {ws_up_r2.max_row}"
empty_count_r2 = 0
invalid_count_r2 = 0
for r in range(2, 280):
    val = ws_up_r2.cell(r, 6).value
    if val is None or str(val).strip() == '':
        empty_count_r2 += 1
    elif str(val).strip() not in valid_funcs:
        invalid_count_r2 += 1

assert empty_count_r2 == 0, f"Found {empty_count_r2} empty cells in Reward (2) Col F!"
assert invalid_count_r2 == 0, f"Found {invalid_count_r2} invalid functions in Reward (2) Col F!"
print("Reward (2) Col F: 278 / 278 rows populated with 100% valid functions.")

diffs_r2 = 0
for r in range(1, 280):
    for c in range(1, ws_bk_r2.max_column + 1):
        if c == 6:
            continue
        v_bk = ws_bk_r2.cell(r, c).value
        v_up = ws_up_r2.cell(r, c).value
        if v_bk != v_up:
            diffs_r2 += 1

assert diffs_r2 == 0, f"Found {diffs_r2} unexpected differences in Reward (2) other columns!"
print("Reward (2) other columns: 100% identical to backup (all formulas and values preserved).")

# Sample prints
sample_rows = [2, 5, 7, 14, 18, 20, 24, 31, 78, 90, 96, 100, 107, 126, 132, 151, 164, 181, 184, 195, 201, 218, 231, 236, 240, 262]
print("\nSample verified rows from MAIN VDC_REWARD:")
for r in sample_rows:
    lc = ws_up_main.cell(r, 2).value
    ld = ws_up_main.cell(r, 3).value
    fn = ws_up_main.cell(r, 6).value
    print(f"  Row {r:3d} | [{lc:8s}] {ld:36s} -> {fn}")

print("\n--- ALL INTEGRITY CHECKS PASSED PERFECTLY ---")
