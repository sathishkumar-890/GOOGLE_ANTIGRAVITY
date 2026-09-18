import openpyxl
from openpyxl.styles import Font, Alignment
import shutil
import os
import sys

sys.path.append('01_VL_EXTRACTION')
from build_reward_function_mapping import get_function_for_reward_row, valid_funcs

EXCEL_PATH = '01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC.xlsx'
BACKUP_PATH = '01_VL_EXTRACTION/VL_EXTRACTION_I.REWARD_MAIN VDC_BACKUP.xlsx'
FUN_PATH = '01_VL_EXTRACTION/C.V. IRENES REWARD_FUN.xlsx'

# 1. Create backup
shutil.copy2(EXCEL_PATH, BACKUP_PATH)
print(f'Backup saved to {BACKUP_PATH}')

# 2. Load workbook
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)

data_font = Font(name='Calibri', size=11)
align = Alignment(horizontal='left', vertical='center')

# 3. Update MAIN VDC_REWARD
ws_main = wb['MAIN VDC_REWARD']
header_f = ws_main.cell(1, 6).value
assert header_f == 'Function Description', f"Expected Col 6 'Function Description', got {header_f}"

main_updated = 0
for r in range(2, ws_main.max_row + 1):
    loc_code = ws_main.cell(r, 2).value
    loc_desc = ws_main.cell(r, 3).value
    sys_desc = ws_main.cell(r, 5).value
    maker = ws_main.cell(r, 7).value
    part = ws_main.cell(r, 10).value
    model = ws_main.cell(r, 16).value

    func, mach = get_function_for_reward_row(r, loc_code, loc_desc, sys_desc, maker, model, part)
    assert func is not None, f"Row {r} could not be mapped: {loc_code} | {loc_desc}"
    assert func in valid_funcs, f"Row {r} function '{func}' not in valid FUN functions!"

    cell = ws_main.cell(row=r, column=6)
    cell.value = func
    cell.font = data_font
    cell.alignment = align
    main_updated += 1

ws_main.column_dimensions['F'].width = 38
print(f"Successfully updated {main_updated} rows in 'MAIN VDC_REWARD' (Col F)")

# 4. Update Reward (2) if present
if 'Reward (2)' in wb.sheetnames:
    ws_r2 = wb['Reward (2)']
    header_r2_f = ws_r2.cell(1, 6).value
    assert header_r2_f == 'FUNCTION DESCRIPTION', f"Expected Col 6 'FUNCTION DESCRIPTION', got {header_r2_f}"

    r2_updated = 0
    for r in range(2, ws_r2.max_row + 1):
        loc_code = ws_r2.cell(r, 2).value
        loc_desc = ws_r2.cell(r, 3).value
        sys_desc = ws_r2.cell(r, 5).value
        maker = ws_r2.cell(r, 7).value
        part = ws_r2.cell(r, 10).value
        model = ws_r2.cell(r, 16).value

        func, mach = get_function_for_reward_row(r, loc_code, loc_desc, sys_desc, maker, model, part)
        assert func is not None, f"Reward (2) Row {r} could not be mapped: {loc_code} | {loc_desc}"
        assert func in valid_funcs, f"Reward (2) Row {r} function '{func}' not in valid FUN functions!"

        cell = ws_r2.cell(row=r, column=6)
        cell.value = func
        cell.font = data_font
        cell.alignment = align
        r2_updated += 1

    ws_r2.column_dimensions['F'].width = 38
    print(f"Successfully updated {r2_updated} rows in 'Reward (2)' (Col F)")

# 5. Save workbook
wb.save(EXCEL_PATH)
print(f"Saved updated workbook to {EXCEL_PATH}")
