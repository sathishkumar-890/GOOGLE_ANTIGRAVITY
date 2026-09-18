import csv
import json
import os
import sys
import tempfile
import time
import requests
import re
from selenium import webdriver
from selenium.webdriver.chrome.options import Options

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# HTML Harness supporting concurrent multi-video playback verification
HTML_POOL = """<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>Live vs Dead Stream Playback Verifier</title>
    <script src="https://cdn.jsdelivr.net/npm/hls.js@latest"></script>
</head>
<body style="background:#0b0e14; color:#fff; font-family:sans-serif;">
    <h3>Active Video Decoder Verification Pool</h3>
    <div id="slots" style="display:flex; gap:12px; margin-bottom:16px;"></div>
    <div id="progress" style="font-size:14px; margin-bottom:8px;">Ready</div>
    <div id="logs" style="font-family:monospace; font-size:11px; height:260px; overflow-y:auto; background:#161b22; padding:8px; border-radius:6px;"></div>
    <script>
        window.testSingleChannel = function(ch, slotIndex, timeoutMs) {
            return new Promise((resolve) => {
                const slots = document.getElementById('slots');
                let slot = document.getElementById('slot_' + slotIndex);
                if (!slot) {
                    slot = document.createElement('div');
                    slot.id = 'slot_' + slotIndex;
                    slot.style.border = '1px solid #30363d';
                    slot.style.padding = '6px';
                    slot.style.borderRadius = '6px';
                    slots.appendChild(slot);
                }
                slot.innerHTML = '<div style="font-size:11px; margin-bottom:4px; max-width:200px; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;"><b>Slot ' + slotIndex + ':</b> #' + ch.row + ' ' + ch.name + '</div>';

                const video = document.createElement('video');
                video.muted = true;
                video.playsInline = true;
                video.autoplay = true;
                video.style.width = '180px';
                video.style.height = '101px';
                video.style.background = '#000';
                slot.appendChild(video);

                let hls = null;
                let resolved = false;
                let fatalReason = null;

                function cleanup(result) {
                    if (resolved) return;
                    resolved = true;
                    clearInterval(checkInterval);
                    clearTimeout(timeoutTimer);
                    if (hls) {
                        try { hls.destroy(); } catch(e) {}
                    }
                    try {
                        video.pause();
                        video.src = '';
                        video.remove();
                    } catch(e) {}
                    resolve(result);
                }

                if (Hls.isSupported()) {
                    hls = new Hls({
                        enableWorker: false,
                        manifestLoadingTimeOut: 6000,
                        manifestLoadingMaxRetry: 1,
                        levelLoadingTimeOut: 6000,
                        levelLoadingMaxRetry: 1,
                        fragLoadingTimeOut: 6000,
                        fragLoadingMaxRetry: 1,
                        startFragPrefetch: true
                    });

                    hls.on(Hls.Events.ERROR, function(event, data) {
                        if (data.fatal) {
                            fatalReason = (data.details || data.type) + (data.response ? (' HTTP ' + data.response.code) : '');
                            cleanup({
                                row: ch.row,
                                name: ch.name,
                                url: ch.url,
                                is_live: false,
                                reason: fatalReason,
                                resolution: '0x0',
                                currentTime: 0
                            });
                        }
                    });

                    hls.on(Hls.Events.MANIFEST_PARSED, function() {
                        video.play().catch(() => {});
                    });

                    hls.loadSource(ch.url);
                    hls.attachMedia(video);
                } else if (video.canPlayType('application/vnd.apple.mpegurl')) {
                    video.src = ch.url;
                    video.addEventListener('loadedmetadata', () => video.play().catch(() => {}));
                } else {
                    cleanup({
                        row: ch.row,
                        name: ch.name,
                        url: ch.url,
                        is_live: false,
                        reason: 'HLS not supported',
                        resolution: '0x0',
                        currentTime: 0
                    });
                    return;
                }

                const startTime = Date.now();
                const checkInterval = setInterval(() => {
                    const elapsed = (Date.now() - startTime) / 1000;
                    const w = video.videoWidth;
                    const h = video.videoHeight;
                    const ct = video.currentTime;

                    // STRICT VIDEO PLAYING VERIFICATION:
                    // 1. Both width and height must be > 0 (decoded video frame)
                    // 2. currentTime must be advancing
                    if (w > 0 && h > 0 && ct > 0.05) {
                        cleanup({
                            row: ch.row,
                            name: ch.name,
                            url: ch.url,
                            is_live: true,
                            reason: 'Video frames actively playing',
                            resolution: w + 'x' + h,
                            currentTime: Math.round(ct * 10) / 10,
                            elapsedSeconds: Math.round(elapsed * 10) / 10
                        });
                    }
                }, 150);

                const timeoutTimer = setTimeout(() => {
                    cleanup({
                        row: ch.row,
                        name: ch.name,
                        url: ch.url,
                        is_live: false,
                        reason: fatalReason || 'Timeout: No video frames decoded within ' + (timeoutMs/1000) + 's',
                        resolution: (video.videoWidth || 0) + 'x' + (video.videoHeight || 0),
                        currentTime: Math.round((video.currentTime || 0) * 10) / 10
                    });
                }, timeoutMs);
            });
        };

        // Worker pool of size N
        window.activeResults = [];
        window.runPool = async function(channelList, concurrency, timeoutMs) {
            window.activeResults = [];
            let nextIndex = 0;
            const prog = document.getElementById('progress');

            async function worker(workerId) {
                while (nextIndex < channelList.length) {
                    const currentIdx = nextIndex++;
                    const ch = channelList[currentIdx];
                    prog.textContent = 'Testing ' + (window.activeResults.length + 1) + ' / ' + channelList.length + '...';
                    const res = await window.testSingleChannel(ch, workerId, timeoutMs);
                    window.activeResults.push(res);
                }
            }

            const workers = [];
            for (let i = 0; i < concurrency; i++) {
                workers.push(worker(i + 1));
            }
            await Promise.all(workers);
            prog.textContent = 'Done! Tested ' + window.activeResults.length + ' channels.';
            return window.activeResults;
        };
    </script>
</body>
</html>
"""

def parse_sheet_csv(csv_path):
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        all_rows = list(reader)

    channels = []
    for idx, r in enumerate(all_rows):
        if not r or len(r) < 2:
            continue
        name = r[0].strip()
        url = r[1].strip()
        if name.upper() == "STREAM NAME" or not url.startswith("http"):
            continue
        cat = r[2].strip() if len(r) > 2 else ""
        icon = r[3].strip() if len(r) > 3 else ""
        status = r[4].strip() if len(r) > 4 else ""
        channels.append({
            "id": f"row_{idx+1}",
            "row": idx + 1,
            "raw_row_idx": idx,
            "name": name,
            "url": url,
            "cat": cat,
            "icon": icon,
            "orig_status": status
        })

    return all_rows, channels

def main():
    csv_path = os.path.abspath("05_STREAMING_PORTAL/sheet2_raw.csv")
    if not os.path.exists(csv_path):
        print(f"Error: {csv_path} not found!")
        sys.exit(1)

    raw_rows, all_channels = parse_sheet_csv(csv_path)
    print(f"Loaded {len(all_channels)} valid stream rows from Sheet2.", flush=True)

    # Deduplicate by URL so we only test each unique stream once
    unique_channels = {}
    for ch in all_channels:
        u = ch["url"]
        if u not in unique_channels:
            unique_channels[u] = ch

    channels = list(unique_channels.values())
    total_channels = len(channels)
    print(f"Identified {total_channels} unique stream URLs to test in Sheet2.", flush=True)

    harness_path = os.path.abspath("05_STREAMING_PORTAL/sheet2_harness.html")
    with open(harness_path, "w", encoding="utf-8") as f:
        f.write(HTML_POOL)

    temp_dir = tempfile.mkdtemp()
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--mute-audio")
    options.add_argument("--disable-web-security")
    options.add_argument("--allow-running-insecure-content")
    options.add_argument(f"--user-data-dir={temp_dir}")
    options.add_argument("--autoplay-policy=no-user-gesture-required")

    driver = webdriver.Chrome(options=options)
    driver.get(f"file:///{harness_path.replace(os.sep, '/')}")
    time.sleep(1)

    driver.set_script_timeout(360)

    CONCURRENCY = 4
    TIMEOUT_MS = 8000
    CHUNK_SIZE = 20

    print(f"\n=======================================================", flush=True)
    print(f"VERIFYING SHEET2 VIDEO PLAYBACK (LIVE vs DEAD)", flush=True)
    print(f"Criteria: Video Width > 0, Height > 0, currentTime > 0.05s", flush=True)
    print(f"Total Unique Streams: {total_channels} | Concurrency: {CONCURRENCY} | Timeout: {TIMEOUT_MS/1000}s", flush=True)
    print(f"=======================================================\n", flush=True)

    start_time = time.time()
    results = []

    for chunk_start in range(0, total_channels, CHUNK_SIZE):
        chunk = channels[chunk_start:chunk_start + CHUNK_SIZE]
        chunk_idx = (chunk_start // CHUNK_SIZE) + 1
        total_chunks = (total_channels + CHUNK_SIZE - 1) // CHUNK_SIZE
        print(f"\n>>> Running Chunk {chunk_idx}/{total_chunks} (Streams {chunk_start+1} to {min(chunk_start+CHUNK_SIZE, total_channels)})...", flush=True)

        chunk_results = driver.execute_async_script(
            """
            const chs = arguments[0];
            const concurrency = arguments[1];
            const timeout = arguments[2];
            const done = arguments[3];
            window.runPool(chs, concurrency, timeout).then(res => done(res));
            """,
            chunk,
            CONCURRENCY,
            TIMEOUT_MS
        )

        for r in chunk_results:
            tag = "[LIVE]" if r["is_live"] else "[DEAD]"
            res_str = f"[{r['resolution']}]" if r["is_live"] else ""
            name_clean = r['name'][:24].encode('ascii', errors='replace').decode('ascii')
            print(f"{tag:6s} {res_str:11s} | {name_clean:24s} | {r['reason']}", flush=True)
            results.append(r)

    total_time = time.time() - start_time
    driver.quit()

    live_count = sum(1 for r in results if r["is_live"])
    dead_count = sum(1 for r in results if not r["is_live"])

    print(f"\n=======================================================", flush=True)
    print(f"VERIFICATION COMPLETED in {total_time:.1f}s ({total_time/60:.1f} min)", flush=True)
    print(f"Total Unique Streams Tested: {len(results)}", flush=True)
    print(f"🟢 LIVE Streams (Playing Video) : {live_count}", flush=True)
    print(f"🔴 DEAD Streams (Not Playing)   : {dead_count}", flush=True)
    print(f"=======================================================\n", flush=True)

    # Save verification JSON
    json_path = os.path.abspath("05_STREAMING_PORTAL/sheet2_live_dead_results.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)
    print(f"Saved JSON verification log: {json_path}", flush=True)

    # Map URL to result
    url_result_map = {r["url"]: r for r in results}

    # Update all rows in Sheet2
    updated_rows = []
    for idx, r in enumerate(raw_rows):
        if idx == 0:
            while len(r) < 5:
                r.append("")
            r[4] = "STATUS"
            updated_rows.append(r[:5])
            continue

        if not r or len(r) < 2 or not r[1].strip().startswith("http"):
            updated_rows.append(r)
            continue

        u = r[1].strip()
        while len(r) < 5:
            r.append("")

        if u in url_result_map:
            res = url_result_map[u]
            r[4] = "Live" if res["is_live"] else "Dead"
        else:
            r[4] = "Dead"
        updated_rows.append(r[:5])

    # Save Updated Sheet2 CSV
    updated_csv_path = os.path.abspath("05_STREAMING_PORTAL/Sheet2_Updated_Status.csv")
    with open(updated_csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(updated_rows)
    print(f"Generated Updated Sheet2 CSV: {updated_csv_path}", flush=True)

    # Save Formatted Excel for Sheet2
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet2 Status"

        header_fill = PatternFill(start_color="1A1D20", end_color="1A1D20", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        live_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
        live_font = Font(name="Calibri", size=10, bold=True, color="155724")

        dead_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        dead_font = Font(name="Calibri", size=10, bold=True, color="721C24")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E0'),
            right=Side(style='thin', color='CBD5E0'),
            top=Side(style='thin', color='CBD5E0'),
            bottom=Side(style='thin', color='CBD5E0')
        )

        for row_idx, r in enumerate(updated_rows, start=1):
            for col_idx, val in enumerate(r, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 5:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if str(val).strip() == "Live":
                        cell.fill = live_fill
                        cell.font = live_font
                    elif str(val).strip() == "Dead":
                        cell.fill = dead_fill
                        cell.font = dead_font

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        excel_path = os.path.abspath("05_STREAMING_PORTAL/Sheet2_Updated_Status.xlsx")
        wb.save(excel_path)
        print(f"Generated Formatted Excel: {excel_path}", flush=True)
    except Exception as e:
        print(f"Excel generation note: {e}", flush=True)

if __name__ == "__main__":
    main()

