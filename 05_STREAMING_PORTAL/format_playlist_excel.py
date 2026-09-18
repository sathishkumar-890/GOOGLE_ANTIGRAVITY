import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

csv_path = "05_STREAMING_PORTAL/IPTV_Playlist_Updated.csv"
xlsx_path = "05_STREAMING_PORTAL/IPTV_Playlist_Updated.xlsx"

with open(csv_path, "r", encoding="utf-8") as f:
    rows = list(csv.reader(f))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "IPTV Verified Playlist"

header_fill = PatternFill(start_color="1A1D20", end_color="1A1D20", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

playing_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
playing_font = Font(name="Calibri", size=10, bold=True, color="155724")

offline_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
offline_font = Font(name="Calibri", size=10, bold=True, color="721C24")

thin_border = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

for row_idx, r in enumerate(rows, start=1):
    for col_idx, val in enumerate(r, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

        if row_idx == 1:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == 5:  # STATUS
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if str(val).strip() == "Playing":
                cell.fill = playing_fill
                cell.font = playing_font
            else:
                cell.fill = offline_fill
                cell.font = offline_font
        elif col_idx == 6:  # RESOLUTION
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Calibri", size=10, bold=True, color="2B6CB0")

for col in ws.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(len(str(cell.value or '')) for cell in col)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(xlsx_path)
print(f"Generated beautifully styled {xlsx_path}")
