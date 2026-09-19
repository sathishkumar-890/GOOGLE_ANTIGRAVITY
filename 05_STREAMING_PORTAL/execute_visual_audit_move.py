import requests
import csv
import io
import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

sys.stdout.reconfigure(encoding="utf-8")

DEAD_NAMES = {
    "MK Six TV1",
    "IBC Tamil",
    "Manorama Tv",
    "National Geographic",
    "News 7 Tv",
    "Sooriyan TV",
    "Star Vijay HD",
    "Suriyan Tv",
    "Vijay Takkar APAC",
    "We Tv"
}

WEBHOOK_URL = "https://script.google.com/macros/s/AKfycbwhOmBGplMtm5BIRP58sG6KOmazzrhmA5LypzMqoA17K_y6m96CteMRQ6sXnCidgaUZ/exec"

def main():
    print("==========================================================================")
    print("MOVING 10 VISUALLY DEAD STREAMS FROM SHEET 1 TO SHEET 2")
    print("==========================================================================")

    # 1. Read current Sheet 1
    s1_path = "05_STREAMING_PORTAL/Sheet1_Merged_Final.csv"
    with open(s1_path, "r", encoding="utf-8") as f:
        s1_all = list(csv.reader(f))
    s1_header = s1_all[0][:5]
    s1_rows = s1_all[1:]

    # 2. Read current Sheet 2
    s2_path = "05_STREAMING_PORTAL/Sheet2_Dead_Only.csv"
    with open(s2_path, "r", encoding="utf-8") as f:
        s2_all = list(csv.reader(f))
    s2_header = s2_all[0][:5]
    s2_rows = s2_all[1:]

    live_s1 = []
    dead_to_move = []

    for r in s1_rows:
        name = r[0].strip()
        if name in DEAD_NAMES:
            r[4] = "Dead"
            dead_to_move.append(r)
        else:
            r[4] = "Live"
            live_s1.append(r)

    print(f"Sheet 1 Live remaining: {len(live_s1)}")
    print(f"Dead streams to move to Sheet 2: {len(dead_to_move)}")

    # Deduplicate when appending to Sheet 2
    existing_s2_urls = set(r[1].strip() for r in s2_rows if len(r) > 1)
    new_s2_rows = list(s2_rows)
    added_count = 0
    for r in dead_to_move:
        if r[1].strip() not in existing_s2_urls:
            new_s2_rows.append(r)
            existing_s2_urls.add(r[1].strip())
            added_count += 1
        else:
            print(f"  URL already in Sheet 2: {r[0]}")

    print(f"Sheet 2 total Dead rows: {len(new_s2_rows)} (added {added_count})")

    final_s1 = [s1_header] + live_s1
    final_s2 = [s2_header] + new_s2_rows

    # Save CSVs
    with open(s1_path, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(final_s1)
    print(f"Saved: {s1_path}")

    with open(s2_path, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(final_s2)
    print(f"Saved: {s2_path}")

    portal_csv = "05_STREAMING_PORTAL/IPTV_Playlist_Merged_Primary.csv"
    with open(portal_csv, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(final_s1)
    print(f"Updated portal primary playlist: {portal_csv}")

    # Save Excel master with text only (NO green fills)
    excel_path = "05_STREAMING_PORTAL/Google_Sheets_Complete_Master.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    text_font = Font(name="Segoe UI", size=10, bold=False, color="000000")

    for title, data in [("IPTV_Playlist", final_s1), ("Sheet2", final_s2)]:
        ws = wb.create_sheet(title=title)
        for r_idx, row in enumerate(data, 1):
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    # Text only - no background highlights!
                    cell.font = text_font
                    if c_idx == 5:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 65
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 10
        ws.column_dimensions["E"].width = 14

    wb.save(excel_path)
    print(f"Saved Excel master (text only, 0 highlights): {excel_path}")

    # Sync to Google Sheets
    print("\nSyncing Sheet 1 (IPTV_Playlist) to Google Sheets...")
    p1 = {"action": "sync_all", "sheetName": "IPTV_Playlist", "rows": final_s1}
    r1 = requests.post(WEBHOOK_URL, json=p1, timeout=60)
    print(f"Sheet 1 Sync Status: {r1.status_code}, Response: {r1.text}")

    print("Syncing Sheet 2 (Sheet2) to Google Sheets...")
    p2 = {"action": "sync_all", "sheetName": "Sheet2", "rows": final_s2}
    r2 = requests.post(WEBHOOK_URL, json=p2, timeout=60)
    print(f"Sheet 2 Sync Status: {r2.status_code}, Response: {r2.text}")

    print("\nSUCCESS: Visual audit and migration completed!")

if __name__ == "__main__":
    main()
