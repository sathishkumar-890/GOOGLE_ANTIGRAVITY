import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
import csv
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

def main():
    print("==========================================================================")
    print("REMOVING GREEN HIGHLIGHT FROM LIVE STATUS IN EXCEL & MASTER DATASETS")
    print("==========================================================================")

    s1_csv = "05_STREAMING_PORTAL/Sheet1_Merged_Final.csv"
    s2_csv = "05_STREAMING_PORTAL/Sheet2_Dead_Only.csv"
    excel_path = "05_STREAMING_PORTAL/Google_Sheets_Complete_Master.xlsx"

    with open(s1_csv, "r", encoding="utf-8") as f:
        final_sheet1 = list(csv.reader(f))

    with open(s2_csv, "r", encoding="utf-8") as f:
        final_sheet2 = list(csv.reader(f))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    # LIVE: NO HIGHLIGHT, clean standard black font, regular weight
    live_font = Font(name="Segoe UI", size=10, bold=False, color="000000")
    
    # DEAD: soft red fill and red font
    dead_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    dead_font = Font(name="Segoe UI", size=10, bold=False, color="B91C1C")

    live_count = 0
    dead_count = 0

    for title, data in [("IPTV_Playlist", final_sheet1), ("Sheet2", final_sheet2)]:
        ws = wb.create_sheet(title=title)
        for r_idx, row in enumerate(data, 1):
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    if c_idx == 5:
                        if val == "Live":
                            # No fill applied!
                            cell.font = live_font
                            live_count += 1
                        else:
                            cell.fill = dead_fill
                            cell.font = dead_font
                            dead_count += 1
                        cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 65
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 10
        ws.column_dimensions["E"].width = 14

    wb.save(excel_path)
    print(f"Saved: {excel_path}")
    print(f"Processed: {live_count} Live cells (highlight removed), {dead_count} Dead cells.")

    # Verify no green fills exist in the entire workbook
    wb_check = openpyxl.load_workbook(excel_path)
    green_found = 0
    for s_name in wb_check.sheetnames:
        sheet = wb_check[s_name]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                    rgb = str(cell.fill.start_color.rgb).upper()
                    if "DCFCE7" in rgb or "00FF00" in rgb or "B7E1CD" in rgb:
                        green_found += 1
                        print(f"WARNING: Green found at {s_name}!{cell.coordinate}")
                if cell.font and cell.font.color and cell.font.color.rgb:
                    rgb = str(cell.font.color.rgb).upper()
                    if "15803D" in rgb:
                        green_found += 1
                        print(f"WARNING: Green font found at {s_name}!{cell.coordinate}")

    if green_found == 0:
        print("VERIFIED: Exactly 0 green highlights exist in Google_Sheets_Complete_Master.xlsx!")
    else:
        print(f"ERROR: {green_found} green instances still remain!")

if __name__ == "__main__":
    main()
